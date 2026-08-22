"""Tests for the Excel Loan & Amortization Schedule Generator. Run: python -m pytest"""

import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from loan import (  # noqa: E402
    monthly_payment,
    build_schedule,
    summarize,
    month_labels,
    write_xlsx,
)


def test_monthly_payment_known_mortgage():
    # 300k @ 6.5% for 30y is a standard textbook figure: 1896.20/mo
    m = monthly_payment(300000, 6.5, 360)
    assert abs(m - 1896.20) < 0.01


def test_zero_interest_is_straight_line():
    assert monthly_payment(1200, 0, 12) == 100.0
    s = summarize(1200, 0, 12)
    assert s["total_interest"] == 0.0
    assert s["n_payments"] == 12


def test_schedule_pays_off_exactly():
    sched = build_schedule(25000, 9, 60)
    assert len(sched) == 60
    assert sched[-1].balance == 0.0  # fully amortized, no residual


def test_principal_plus_interest_equals_payment():
    sched = build_schedule(10000, 12, 24)
    for row in sched:
        assert abs(row.payment - (row.principal + row.interest)) < 0.02


def test_total_principal_equals_loan():
    sched = build_schedule(25000, 9, 60)
    total_principal = sum(r.principal for r in sched)
    assert abs(total_principal - 25000) < 0.02


def test_extra_payment_saves_interest_and_time():
    base = summarize(25000, 9, 60)
    with_extra = summarize(25000, 9, 60, extra=100)
    assert with_extra["n_payments"] < base["n_payments"]
    assert with_extra["total_interest"] < base["total_interest"]
    assert with_extra["interest_saved"] > 0
    assert with_extra["months_saved"] == base["n_payments"] - with_extra["n_payments"]


def test_month_labels_wrap_year():
    labels = month_labels("2026-11", 4)
    assert labels == ["2026-11", "2026-12", "2027-01", "2027-02"]


def test_write_xlsx_has_both_sheets(tmp_path):
    from openpyxl import load_workbook

    s = summarize(300000, 6.5, 360)
    out = tmp_path / "sched.xlsx"
    write_xlsx(s, str(out), start_date="2026-09")
    wb = load_workbook(out)
    assert wb.sheetnames == ["Summary", "Schedule"]
    sh = wb["Schedule"]
    assert sh["A1"].value == "Month"  # dated schedule -> month column first
    assert sh.freeze_panes == "A2"
    assert sh.max_row == 361  # header + 360 payments
